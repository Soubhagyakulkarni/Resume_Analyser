"""Text extraction for DOCX resumes."""

from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError

from app.utils.errors import InvalidResumeError
from models.resume import ParsedResume


def parse_docx(filename: str, content: bytes) -> ParsedResume:
    try:
        document = Document(BytesIO(content))
        lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in lines:
                        lines.append(cell_text)
        text = "\n".join(lines).strip()
    except (PackageNotFoundError, XMLSyntaxError, ValueError, KeyError, OSError) as exc:
        raise InvalidResumeError("The DOCX file could not be parsed") from exc
    if not text:
        raise InvalidResumeError("No readable text was found in the DOCX file")
    return ParsedResume(filename=filename, file_type="docx", text=text)
